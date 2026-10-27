#!/usr/bin/env python3
"""Dựng báo cáo chiến lược mẫu (DOCX → PDF) theo layout Promete.

    python build_report.py report.json --out-dir OUT [--no-pdf]

Đầu vào là một file JSON nội dung (xem ../references/layout-spec.md). Bộ dựng
chỉ lo trình bày: bìa, mục lục có số trang, tiêu đề mục, bảng, biểu đồ, ghi
chú. Số liệu và câu chữ phải đến từ nguồn đã kiểm, không sinh ở đây.

Mục lục được dựng hai lượt: lượt đầu xuất PDF để đọc trang thật của từng tiêu
đề, lượt hai ghi số trang vào mục lục (vẫn là trường TOC của Word, bấm F9 để
cập nhật khi sửa DOCX).
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SKILL_DIR = Path(__file__).resolve().parent.parent
ASSETS = SKILL_DIR / "assets"

ORANGE = "E69138"
CREAM = "FFFBEC"
SHADE = "F2F3F5"
RULE = "CCCCCC"
GREY = "999999"
BLACK = "000000"

BODY_FONT = "Lexend Light"
BOLD_FONT = "Lexend"
BODY_PT = 10

PAGE_W_CM, MARGIN_LR_CM = 21.0, 1.5
TEXT_W_CM = PAGE_W_CM - 2 * MARGIN_LR_CM


# ---------------------------------------------------------------- fonts
def ensure_fonts() -> None:
    """Cài Lexend Light/Bold cho người dùng nếu fontconfig chưa thấy."""
    have = subprocess.run(["fc-list", ":family=Lexend Light"], capture_output=True, text=True).stdout
    bold = subprocess.run(["fc-list", ":family=Lexend:style=Bold"], capture_output=True, text=True).stdout
    if have.strip() and bold.strip():
        return
    dest = Path.home() / ".local/share/fonts/promete-lexend"
    dest.mkdir(parents=True, exist_ok=True)
    for f in (ASSETS / "fonts").glob("*.ttf"):
        shutil.copy2(f, dest / f.name)
    subprocess.run(["fc-cache", "-f", str(dest)], check=True, capture_output=True)


# ---------------------------------------------------------------- xml helpers
def _el(tag: str, **attrs) -> OxmlElement:
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_run_font(run, bold: bool = False, italic: bool = False, size: float | None = None,
                 color: str | None = None) -> None:
    name = BOLD_FONT if bold else BODY_FONT
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = _el("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(k), name)
    run.bold = bold or None
    run.italic = italic or None
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def para_border(paragraph, side: str, color: str, sz: int = 8, space: int = 4, val: str = "single") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = _el("w:pBdr")
        ppr.append(pbdr)
    pbdr.append(_el(f"w:{side}", **{"w:val": val, "w:sz": sz, "w:space": space, "w:color": color}))


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcpr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))


def cell_borders(cell, **sides) -> None:
    """sides: top=('single', 8, 'E69138') ... ; None → nil."""
    tcpr = cell._tc.get_or_add_tcPr()
    b = tcpr.find(qn("w:tcBorders"))
    if b is None:
        b = _el("w:tcBorders")
        tcpr.append(b)
    for side in ("top", "left", "bottom", "right"):
        if side not in sides:
            continue
        spec = sides[side]
        old = b.find(qn(f"w:{side}"))
        if old is not None:
            b.remove(old)
        if spec is None:
            b.append(_el(f"w:{side}", **{"w:val": "nil"}))
        else:
            val, sz, color = spec
            b.append(_el(f"w:{side}", **{"w:val": val, "w:sz": sz, "w:space": 0, "w:color": color}))


def cell_margins(table, top=75, bottom=75, left=110, right=110) -> None:
    tblpr = table._tbl.tblPr
    m = _el("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        m.append(_el(f"w:{side}", **{"w:w": v, "w:type": "dxa"}))
    tblpr.append(m)


def no_table_borders(table) -> None:
    tblpr = table._tbl.tblPr
    b = _el("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(_el(f"w:{side}", **{"w:val": "nil"}))
    tblpr.append(b)


def fixed_widths(table, widths_cm: list[float]) -> None:
    tblpr = table._tbl.tblPr
    tblpr.append(_el("w:tblLayout", **{"w:type": "fixed"}))
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = _el("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tblw.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def row_flags(row, header: bool = False) -> None:
    trpr = row._tr.get_or_add_trPr()
    trpr.append(_el("w:cantSplit"))
    if header:
        trpr.append(_el("w:tblHeader"))


def add_bookmark(paragraph, name: str, bid: int) -> None:
    start = _el("w:bookmarkStart", **{"w:id": bid, "w:name": name})
    end = _el("w:bookmarkEnd", **{"w:id": bid})
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


# ---------------------------------------------------------------- inline markup
_TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)")


def add_rich(paragraph, text: str, size: float | None = None, color: str | None = None,
             bold_all: bool = False, italic_all: bool = False) -> None:
    """**đậm**, *nghiêng*, xuống dòng bằng \\n."""
    lines = str(text).split("\n")
    for li, line in enumerate(lines):
        for part in _TOKEN.split(line):
            if not part:
                continue
            bold, italic = bold_all, italic_all
            if part.startswith("**") and part.endswith("**"):
                part, bold = part[2:-2], True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                part, italic = part[1:-1], True
            run = paragraph.add_run(part)
            set_run_font(run, bold=bold, italic=italic, size=size, color=color)
        if li < len(lines) - 1:
            br = paragraph.add_run()
            br.add_break()


def fmt(paragraph, *, align=None, before=0, after=None, line=None, keep_next=False,
        left_cm=None, hanging_cm=None, page_break_before=False) -> None:
    pf = paragraph.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if keep_next:
        pf.keep_with_next = True
    if left_cm is not None:
        pf.left_indent = Cm(left_cm)
    if hanging_cm is not None:
        pf.first_line_indent = Cm(-hanging_cm)
    if page_break_before:
        pf.page_break_before = True


# ---------------------------------------------------------------- document setup
def setup_document(opts: dict) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(MARGIN_LR_CM)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(1.8)
    sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_PT)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(k), BODY_FONT)
    lang = _el("w:lang", **{"w:val": "vi-VN", "w:eastAsia": "vi-VN", "w:bidi": "ar-SA"})
    rpr.append(lang)
    npf = normal.paragraph_format
    npf.space_after = Pt(8)
    npf.line_spacing = 1.15
    npf.widow_control = True

    for name, size, color in (("Heading 1", 13, ORANGE), ("Heading 2", 10.5, BLACK)):
        st = doc.styles[name]
        st.font.name = BOLD_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor.from_string(color)
        srpr = st.element.get_or_add_rPr()
        f = srpr.find(qn("w:rFonts"))
        if f is None:
            f = _el("w:rFonts")
            srpr.insert(0, f)
        for a in list(f.attrib):
            del f.attrib[a]
        for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            f.set(qn(k), BOLD_FONT)
        st.paragraph_format.keep_with_next = True

    # logo góc phải trên mọi trang, kể cả bìa
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.add_run().add_picture(str(ASSETS / "promete-logo.jpg"), width=Cm(float(opts.get("logo_width_cm", 3.1))))

    if opts.get("page_numbers"):
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_field(fp, "PAGE", "1", size=8, color=GREY)
    return doc


def add_field(paragraph, instr: str, cached: str, size=None, color=None) -> None:
    r = paragraph.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    r = paragraph.add_run()
    it = _el("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    r = paragraph.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    r = paragraph.add_run(cached)
    set_run_font(r, size=size, color=color)
    r = paragraph.add_run()
    r._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


# ---------------------------------------------------------------- blocks
class Builder:
    def __init__(self, spec: dict, base: Path, pages: dict[str, int] | None):
        self.spec = spec
        self.base = base
        self.pages = pages or {}
        self.opts = spec.get("options", {})
        self.doc = setup_document(self.opts)
        self.figure_no = 0
        self.bookmark_id = 100
        self.headings: list[dict] = []  # {text, level, anchor}
        self.tmp = Path(tempfile.mkdtemp(prefix="promete-report-"))

    # -- numbering
    def plan_headings(self) -> None:
        n = 0
        for si, sec in enumerate(self.spec["sections"]):
            numbered = sec.get("numbered", True)
            if numbered:
                n += 1
            sec["_label"] = f"{n}. {sec['title']}" if numbered else sec["title"]
            sec["_anchor"] = f"_Toc_s{si}"
            self.headings.append({"text": sec["_label"], "level": 1, "anchor": sec["_anchor"]})
            m = 0
            for bi, blk in enumerate(sec.get("blocks", [])):
                if blk.get("type") == "h2":
                    m += 1
                    blk["_label"] = f"{n}.{m}. {blk['text']}" if numbered else blk["text"]
                    blk["_anchor"] = f"_Toc_s{si}_{bi}"
                    self.headings.append({"text": blk["_label"], "level": 2, "anchor": blk["_anchor"]})

    # -- cover
    def cover(self) -> None:
        c = self.spec["cover"]
        doc = self.doc
        p = doc.add_paragraph()
        fmt(p, before=int(self.opts.get("cover_top_pt", 70)), after=14)
        add_rich(p, c.get("kicker", "Báo cáo chiến lược mẫu"), size=12, color=GREY)
        p = doc.add_paragraph()
        fmt(p, after=10)
        add_rich(p, c["number_line"], size=16, bold_all=True, color=BLACK)
        p = doc.add_paragraph()
        fmt(p, after=48, line=1.05)
        add_rich(p, c["title"], size=26, bold_all=True, color=ORANGE)

        facts = c.get("facts", [])
        if facts:
            t = doc.add_table(rows=len(facts), cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.LEFT
            cell_margins(t, top=70, bottom=70, left=120, right=120)
            fixed_widths(t, [5.2, TEXT_W_CM - 5.2])
            for r, (k, v) in enumerate(facts):
                row = t.rows[r]
                row_flags(row)
                for ci, txt in enumerate((k, v)):
                    cell = row.cells[ci]
                    cell_borders(cell, **{s: ("single", 6, RULE) for s in ("top", "left", "bottom", "right")})
                    if ci == 0:
                        shade_cell(cell, CREAM)
                    cp = cell.paragraphs[0]
                    fmt(cp, after=0, line=1.15)
                    add_rich(cp, txt, size=BODY_PT)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # -- toc
    def toc(self) -> None:
        tcfg = self.spec.get("toc", {})
        doc = self.doc
        p = doc.add_paragraph()
        fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
        add_rich(p, tcfg.get("title", "Mục lục"), size=13, bold_all=True, color=ORANGE)
        entries = self.headings
        for i, h in enumerate(entries):
            ep = doc.add_paragraph()
            indent = 0.8 if h["level"] == 2 else 0
            fmt(ep, after=3, line=1.15, left_cm=indent)
            ep.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if i == 0:
                r = ep.add_run()
                r._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
                r = ep.add_run()
                it = _el("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = ' TOC \\o "1-2" \\h \\z \\u '
                r._r.append(it)
                r = ep.add_run()
                r._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
            link = _el("w:hyperlink", **{"w:anchor": h["anchor"], "w:history": 1})
            page = str(self.pages.get(h["anchor"], 0))
            for txt in (h["text"], "\t", page):
                r = ep.add_run(txt)
                set_run_font(r, size=BODY_PT)
                link.append(r._r)
            ep._p.append(link)
            if i == len(entries) - 1:
                r = ep.add_run()
                r._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # -- sections
    def section(self, sec: dict, first: bool) -> None:
        h = self.doc.add_paragraph(style="Heading 1")
        fmt(h, before=0 if first or sec.get("new_page") else 16, after=12, keep_next=True,
            page_break_before=bool(sec.get("new_page")) and not first)
        para_border(h, "bottom", ORANGE, sz=8, space=6)
        r = h.add_run(sec["_label"])
        set_run_font(r, bold=True)
        self.bookmark_id += 1
        add_bookmark(h, sec["_anchor"], self.bookmark_id)
        for blk in sec.get("blocks", []):
            getattr(self, "b_" + blk["type"])(blk)

    def b_p(self, blk: dict) -> None:
        p = self.doc.add_paragraph()
        fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=blk.get("after", 8), keep_next=blk.get("keep_next", False))
        add_rich(p, blk["text"])

    def b_h2(self, blk: dict) -> None:
        h = self.doc.add_paragraph(style="Heading 2")
        fmt(h, before=10, after=8, keep_next=True)
        r = h.add_run(blk["_label"])
        set_run_font(r, bold=True, size=10.5)
        self.bookmark_id += 1
        add_bookmark(h, blk["_anchor"], self.bookmark_id)

    def _list(self, items: list[str], marker) -> None:
        for i, it in enumerate(items):
            p = self.doc.add_paragraph()
            left, hang = (1.3, 0.55) if marker == "num" else (1.0, 0.5)
            fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6 if marker == "num" else 2,
                left_cm=left, hanging_cm=hang)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(left))
            m = f"{i + 1}." if marker == "num" else "-"
            r = p.add_run(m + "\t")
            set_run_font(r)
            add_rich(p, it)
        # khoảng cách sau cả danh sách
        p.paragraph_format.space_after = Pt(10)

    def b_bullets(self, blk: dict) -> None:
        self._list(blk["items"], "dash")

    def b_numbered(self, blk: dict) -> None:
        self._list(blk["items"], "num")

    def b_note(self, blk: dict) -> None:
        p = self.doc.add_paragraph()
        fmt(p, align=WD_ALIGN_PARAGRAPH.RIGHT, before=2, after=10)
        add_rich(p, blk["text"], italic_all=True, size=9.5)

    def b_page_break(self, blk: dict) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def b_table(self, blk: dict) -> None:
        cols, rows = blk["columns"], blk["rows"]
        n = len(cols)
        fr = blk.get("widths") or [1] * n
        widths = [TEXT_W_CM * f / sum(fr) for f in fr]
        align = blk.get("align") or ["left"] * n
        amap = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                "right": WD_ALIGN_PARAGRAPH.RIGHT}
        shade_first = blk.get("first_col_shade", True)
        rules = blk.get("rules", "none")
        size = blk.get("font_pt", BODY_PT)

        keep_whole = blk.get("keep_together", len(rows) <= 14)
        t = self.doc.add_table(rows=1 + len(rows), cols=n)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        no_table_borders(t)
        cell_margins(t)
        fixed_widths(t, widths)
        for ri, row in enumerate(t.rows):
            header = ri == 0
            row_flags(row, header=header)
            values = cols if header else rows[ri - 1]
            if len(values) != n:
                raise ValueError(f"bảng '{cols[0]}': hàng {ri} có {len(values)} ô, cần {n}")
            last = ri == len(rows)
            for ci, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                b = {"left": None, "right": None}
                if header:
                    shade_cell(cell, CREAM)
                    b["top"] = ("single", 8, ORANGE)
                    b["bottom"] = ("single", 8, ORANGE)
                else:
                    if shade_first and ci == 0:
                        shade_cell(cell, SHADE)
                    if last:
                        b["bottom"] = ("single", 8, RULE)
                    elif rules == "dotted":
                        b["bottom"] = ("dotted", 4, RULE)
                    else:
                        b["bottom"] = None
                cell_borders(cell, **b)
                cp = cell.paragraphs[0]
                fmt(cp, align=amap[align[ci]], after=0, line=1.15)
                txt = values[ci]
                if header:
                    add_rich(cp, txt, size=size, color=ORANGE)
                else:
                    add_rich(cp, txt, size=size)
                # bảng nhỏ giữ trọn một trang: mọi hàng trừ hàng cuối "đi cùng" hàng sau
                cp.paragraph_format.keep_with_next = keep_whole and not last
        if blk.get("note"):
            self.b_note({"text": blk["note"]})
        else:
            sp = self.doc.add_paragraph()
            fmt(sp, after=0, line=0.8)

    def _caption(self, text: str | None) -> None:
        if not text:
            return
        self.figure_no += 1
        p = self.doc.add_paragraph()
        fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=12)
        add_rich(p, f"Hình {self.figure_no}. {text}", size=BODY_PT)

    def b_image(self, blk: dict) -> None:
        p = self.doc.add_paragraph()
        fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2, keep_next=True)
        path = (self.base / blk["path"]).resolve()
        p.add_run().add_picture(str(path), width=Cm(float(blk.get("width_cm", 16))))
        self._caption(blk.get("caption"))

    def b_chart(self, blk: dict) -> None:
        png = render_chart(blk, self.tmp / f"chart{self.figure_no + 1}.png", self.base)
        self.b_image({"path": str(png), "width_cm": blk.get("width_cm", 15.5), "caption": blk.get("caption")})

    def build(self) -> Document:
        self.plan_headings()
        self.cover()
        if self.spec.get("toc", {"enabled": True}).get("enabled", True):
            self.toc()
        for i, sec in enumerate(self.spec["sections"]):
            self.section(sec, first=(i == 0))
        props = self.doc.core_properties
        meta = self.spec.get("meta", {})
        props.title = meta.get("title", self.spec["cover"]["title"])
        props.author = meta.get("author", "Promete")
        props.language = "vi-VN"
        return self.doc


# ---------------------------------------------------------------- charts
def vi_number(x: float, fmt_spec: str) -> str:
    s = fmt_spec.format(x)
    s = s.replace(",", " ").replace(".", ",").replace(" ", ".")
    return s.replace("-", "−")


SERIES_COLORS = {"orange": "#" + ORANGE, "grey": "#" + GREY, "dark": "#434343", "light": "#CCCCCC"}


def _style_axes(ax, matplotlib, ylabel: str | None, yfmt: str = "{:g}") -> None:
    ax.axhline(0, color="#000000", linewidth=0.8, zorder=3)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.spines["left"].set_color("#000000")
    ax.spines["bottom"].set_visible(False)
    ax.tick_params(axis="x", length=0, pad=8, labelsize=12)
    ax.tick_params(axis="y", labelsize=11)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: vi_number(v, yfmt)))
    if ylabel:
        ax.set_ylabel(ylabel, fontsize=12)


def _load_series(s: dict, base: Path) -> tuple[list, list[float]]:
    """Series lấy thẳng từ CSV nguồn: {"csv", "x", "y", "scale", "rebase"} — không gõ số tay."""
    if "csv" not in s:
        return s.get("x", []), s["y"]
    import csv as _csv
    rows = list(_csv.DictReader(open((base / s["csv"]).resolve(), encoding="utf-8")))
    xs, ys = [], []
    for r in rows:
        if r.get(s["y"]) in (None, ""):
            continue
        xs.append(r[s["x"]])
        ys.append(float(r[s["y"]]))
    if s.get("rebase"):  # chuỗi mức (equity/NAV) → lợi suất lũy kế so với điểm đầu
        ys = [y / ys[0] - 1 for y in ys]
    return xs, [y * float(s.get("scale", 1)) for y in ys]


def render_chart(blk: dict, out: Path, base: Path) -> Path:
    import datetime as dt
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.dates as mdates
    import matplotlib.pyplot as plt
    from matplotlib import font_manager

    for f in (ASSETS / "fonts").glob("*.ttf"):
        font_manager.fontManager.addfont(str(f))
    plt.rcParams.update({"font.family": "Lexend", "font.weight": "light", "font.size": 12,
                         "axes.unicode_minus": True})
    kind = blk.get("kind", "bar")
    fig, ax = plt.subplots(figsize=(blk.get("w_in", 7.6), blk.get("h_in", 3.4)), dpi=220)
    vf = blk.get("value_format", "{:.2f}%")

    if kind == "bar" and "series" not in blk:
        labels, values = blk["labels"], blk["values"]
        hi = blk.get("highlight", [0])
        hi = [hi] if isinstance(hi, int) else hi
        colors = ["#" + (ORANGE if i in hi else GREY) for i in range(len(values))]
        bars = ax.bar(labels, values, color=colors, width=0.55, zorder=2)
        _style_axes(ax, matplotlib, blk.get("ylabel"))
        ax.tick_params(axis="x", labelsize=13)
        span = (max(values) - min(min(values), 0)) or 1
        for b, v in zip(bars, values):
            y = v + span * 0.03 if v >= 0 else v - span * 0.03
            ax.text(b.get_x() + b.get_width() / 2, y, vi_number(v, vf), ha="center",
                    va="bottom" if v >= 0 else "top", fontsize=13)
        lo = min(min(values), 0)
        ax.set_ylim(lo - span * 0.14 if lo < 0 else 0, max(max(values), 0) + span * 0.14)
    elif kind == "bar":  # cột nhóm: labels × series
        labels, series = blk["labels"], blk["series"]
        n = len(series)
        w = 0.8 / n
        allv = [v for s in series for v in s["values"]]
        span = (max(allv) - min(min(allv), 0)) or 1
        for k, s in enumerate(series):
            xs = [i - 0.4 + w * (k + 0.5) for i in range(len(labels))]
            bars = ax.bar(xs, s["values"], width=w * 0.92, color=SERIES_COLORS[s.get("color", "grey")],
                          label=s["name"], zorder=2)
            if blk.get("value_labels", True):
                for b, v in zip(bars, s["values"]):
                    y = v + span * 0.02 if v >= 0 else v - span * 0.02
                    ax.text(b.get_x() + b.get_width() / 2, y, vi_number(v, blk.get("value_format", "{:.0f}")),
                            ha="center", va="bottom" if v >= 0 else "top", fontsize=8.5)
        ax.set_xticks(range(len(labels)), labels)
        _style_axes(ax, matplotlib, blk.get("ylabel"))
        ax.set_ylim(min(min(allv), 0) - span * 0.12, max(max(allv), 0) + span * 0.12)
        ax.legend(frameon=False, fontsize=10, ncol=n, loc="upper center", bbox_to_anchor=(0.5, 1.13))
    elif kind == "line":
        for s in blk["series"]:
            xs, ys = _load_series(s, base)
            xs = [dt.date.fromisoformat(x[:10]) for x in xs]
            ax.plot(xs, ys, color=SERIES_COLORS[s.get("color", "grey")], linewidth=s.get("width", 1.6),
                    label=s["name"], zorder=3 if s.get("color") == "orange" else 2)
            if s.get("end_label", True):
                ax.annotate(vi_number(ys[-1], vf), (xs[-1], ys[-1]), xytext=(4, 0), textcoords="offset points",
                            va="center", fontsize=10, color=SERIES_COLORS[s.get("color", "grey")])
        _style_axes(ax, matplotlib, blk.get("ylabel"))
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%m/%Y"))
        ax.xaxis.set_major_locator(mdates.MonthLocator(interval=blk.get("month_step", 2)))
        ax.margins(x=0.08)
        ax.legend(frameon=False, fontsize=10, ncol=len(blk["series"]), loc="upper center",
                  bbox_to_anchor=(0.5, 1.13))
    else:
        raise ValueError("chart.kind: 'bar' (một chuỗi hoặc series) hoặc 'line'")
    fig.tight_layout()
    fig.savefig(out, dpi=220, facecolor="white")
    plt.close(fig)
    return out


# ---------------------------------------------------------------- lint
def lint(spec: dict) -> list[str]:
    """Lỗi trình bày hay gặp ở bản mẫu: số thập phân kiểu Anh, ô trống, placeholder."""
    issues: list[str] = []

    def walk(obj, where):
        if isinstance(obj, str):
            for m in re.finditer(r"(?<![\w.])[+−-]?\d+\.\d+\s?(%|đpt|năm)", obj):
                issues.append(f"{where}: số thập phân dùng dấu chấm, cần dấu phẩy kiểu Việt: {m.group(0)!r}")
            for m in re.finditer(r"TODO|TBD|\?\?\?|<placeholder>|XXX", obj):
                issues.append(f"{where}: còn placeholder: {obj[max(0, m.start() - 30):m.end() + 10]!r}")
        elif isinstance(obj, dict):
            for k, v in obj.items():
                if not k.startswith("_") and k not in ("path", "value_format", "csv", "x", "y"):
                    walk(v, f"{where}.{k}")
        elif isinstance(obj, list):
            for i, v in enumerate(obj):
                walk(v, f"{where}[{i}]")

    walk(spec, "$")
    try:  # ký tự Lexend không có sẽ bị thay bằng font khác trong PDF
        from fontTools.ttLib import TTFont
        cmap = TTFont(str(ASSETS / "fonts" / "Lexend-Light.ttf")).getBestCmap()
        text = json.dumps(spec, ensure_ascii=False)
        missing = sorted({ch for ch in text if ord(ch) > 127 and ord(ch) not in cmap})
        if missing:
            issues.append(f"ký tự không có trong Lexend (PDF sẽ đổi font): {' '.join(missing)}")
    except ImportError:
        pass
    for si, sec in enumerate(spec.get("sections", [])):
        for bi, blk in enumerate(sec.get("blocks", [])):
            if blk.get("type") == "table":
                for ri, row in enumerate(blk["rows"]):
                    for ci, v in enumerate(row):
                        if not str(v).strip():
                            issues.append(f"sections[{si}].blocks[{bi}] hàng {ri + 1} cột {ci + 1}: ô trống"
                                          " → ghi 'chưa đo'/'không áp dụng' hoặc gộp hàng")
    return issues


# ---------------------------------------------------------------- pdf + page numbers
def to_pdf(docx_path: Path, out_dir: Path) -> Path:
    with tempfile.TemporaryDirectory(prefix="lo-profile-") as prof:
        subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless",
                        "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                       check=True, capture_output=True, timeout=240)
    pdf = out_dir / (docx_path.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError("LibreOffice không tạo được PDF")
    return pdf


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", s)).strip().casefold()


def find_pages(pdf: Path, headings: list[dict], start_page: int) -> dict[str, int]:
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", str(pdf)], capture_output=True,
                                                          text=True).stdout).group(1))
    texts = [_norm(subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), str(pdf), "-"],
                                  capture_output=True, text=True).stdout) for i in range(1, n + 1)]
    pages, cur = {}, start_page
    for h in headings:
        needle = _norm(h["text"])
        for p in range(cur, n + 1):
            if needle in texts[p - 1]:
                pages[h["anchor"]] = cur = p
                break
        else:
            raise RuntimeError(f"không tìm thấy tiêu đề trong PDF: {h['text']!r}")
    return pages


def check_fonts(pdf: Path) -> list[str]:
    out = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True).stdout
    names = [ln.split()[0] for ln in out.splitlines()[2:] if ln.strip()]
    bad = [n for n in names if "Lexend" not in n]
    return bad


def build(spec_path: Path, out_dir: Path, make_pdf: bool = True) -> dict:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    base = spec_path.parent
    stem = spec.get("output_name") or spec_path.stem
    out_dir.mkdir(parents=True, exist_ok=True)
    issues = lint(spec)
    ensure_fonts()

    def render(pages):
        b = Builder(copy.deepcopy(spec), base, pages)
        doc = b.build()
        path = out_dir / f"{stem}.docx"
        doc.save(path)
        return b, path

    b, docx_path = render(None)
    result = {"docx": str(docx_path), "lint": issues}
    if not make_pdf:
        return result
    has_toc = spec.get("toc", {"enabled": True}).get("enabled", True)
    with tempfile.TemporaryDirectory(prefix="promete-pass-") as tmp:
        pdf = to_pdf(docx_path, Path(tmp))
        if has_toc:
            pages = find_pages(pdf, b.headings, start_page=3)
            b, docx_path = render(pages)
            pdf = to_pdf(docx_path, Path(tmp))
            again = find_pages(pdf, b.headings, start_page=3)
            if again != pages:
                raise RuntimeError(f"số trang mục lục không ổn định: {pages} → {again}")
            result["toc_pages"] = {h["text"]: pages[h["anchor"]] for h in b.headings}
        final = out_dir / f"{stem}.pdf"
        shutil.copy2(pdf, final)
    result["pdf"] = str(final)
    result["non_lexend_fonts"] = check_fonts(final)
    return result


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", type=Path)
    ap.add_argument("--out-dir", type=Path, required=True)
    ap.add_argument("--no-pdf", action="store_true")
    ap.add_argument("--strict", action="store_true", help="dừng nếu lint có cảnh báo")
    a = ap.parse_args()
    res = build(a.spec, a.out_dir, make_pdf=not a.no_pdf)
    print(json.dumps(res, ensure_ascii=False, indent=2))
    if a.strict and res["lint"]:
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
